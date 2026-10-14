"""报价估算书（项目制）排成 Word。数据来自 quotes 表；对比常规外包、优势、付款、资源费都按剑来口径生成。"""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.contract_doc import PARTY_B, cn_date, money, pct, rmb_upper
from app.timezone import now_cn

# 常规外包单价区间（元 / 人天）与人天放大系数：多角色交接、会议文档成本
MARKET_RATE = (1000, 1200)
MARKET_DAYS = (1.3, 1.4)
MARKET_WEEKS = (1.4, 1.7)


def quote_totals(items: list[dict], ai_fee: Decimal) -> dict:
    dev = sum((Decimal(str(i.get("amount") or 0)) for i in items), Decimal(0))
    days = sum((Decimal(str(i.get("days") or 0)) for i in items), Decimal(0))
    return {"dev": dev, "days": days, "ai_fee": ai_fee, "total": dev + ai_fee}


def _round_wan(value: float) -> str:
    return f"{value / 10000:.0f}"


class _Doc:
    def __init__(self) -> None:
        self.doc = Document()
        for s in self.doc.sections:
            s.top_margin = s.bottom_margin = Cm(2.2)
            s.left_margin = s.right_margin = Cm(2.4)

    @staticmethod
    def font(run, size=10.5, bold=False, gray=False):
        run.font.name = "宋体"
        run.font.size = Pt(size)
        run.font.bold = bold
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        if gray:
            run.font.color.rgb = RGBColor(90, 90, 90)

    def para(self, text="", size=10.5, bold=False, align=None, indent=False, gray=False):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.4
        if indent:
            p.paragraph_format.first_line_indent = Pt(size * 2)
        if align is not None:
            p.alignment = align
        if text:
            self.font(p.add_run(text), size, bold, gray)
        return p

    def h1(self, text):
        self.para(text, 14, True)

    def h2(self, text):
        self.para(text, 11.5, True)

    def bullet(self, text):
        self.para("· " + text).paragraph_format.left_indent = Pt(12)

    def table(self, head, rows, widths, bold_last=False):
        t = self.doc.add_table(rows=1, cols=len(head))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for c, txt in zip(t.rows[0].cells, head):
            self.font(c.paragraphs[0].add_run(txt), 10, True)
        for i, r in enumerate(rows):
            cells = t.add_row().cells
            last = bold_last and i == len(rows) - 1
            for c, txt in zip(cells, r):
                self.font(c.paragraphs[0].add_run(str(txt)), 10, last)
        for row in t.rows:
            for c, w in zip(row.cells, widths):
                c.width = Cm(w)
        self.para()


def build_quote_docx(q: dict) -> bytes:
    items: list[dict] = q["items"]
    pay_plan: list[dict] = q["pay_plan"]
    schedule: list[dict] = q["schedule"]
    ai_fee = Decimal(str(q["ai_fee"]))
    sums = quote_totals(items, ai_fee)
    total, dev, days = sums["total"], sums["dev"], sums["days"]
    weeks = q["weeks"]
    onsite = q["onsite_per_week"]
    m_low, m_high = money(Decimal(str(q["maintain_low"]))), money(Decimal(str(q["maintain_high"])))
    made = q.get("created_on") or now_cn().date()

    d = _Doc()
    d.para(q["title"], 18, True, WD_ALIGN_PARAGRAPH.CENTER)
    d.para("报价估算书", 16, True, WD_ALIGN_PARAGRAPH.CENTER)
    d.para(
        f"报价编号：{q['code']}　　致：{q['party_a']}　　服务方：{PARTY_B['name']}"
        f"　　编制日期：{cn_date(made)}　　有效期：{q['valid_days']} 日",
        9.5,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        gray=True,
    )
    d.para()

    d.h1("一、结论")
    summary = [
        [
            "总价",
            f"人民币 {money(total)} 元（大写：{rmb_upper(total)}），含税，固定总价"
            + (
                f"；含开发服务费 {money(dev)} 元 + 开发期 AI 费用 {money(ai_fee)} 元（包干）"
                if ai_fee
                else ""
            ),
        ],
        ["工期", f"{weeks} 周（自预付款到账、需求清单确认之日起算）"],
        ["交付", "完整源码 + 部署上线 + 使用培训 + 1 个月质保；源码及设计版权付清后归甲方"],
    ]
    if onsite:
        summary.append(
            [
                "驻场沟通",
                f"开发期每周 {onsite} 次到甲方现场沟通（演示进度、确认需求、当场调整），含在总价内",
            ]
        )
    if pay_plan:
        summary.append(
            [
                "付款",
                " / ".join(f"{pct(p['pct'])}%" for p in pay_plan)
                + f" 共 {len(pay_plan)} 期，按里程碑支付",
            ]
        )
    summary.append(
        [
            "运维",
            f"上线后可选月度运维 {m_low}–{m_high} 元 / 月，只含缺陷修复与小需求；小程序、App 等大功能另签项目",
        ]
    )
    summary.append(
        [
            "不含",
            "项目运行所需服务器、短信、第三方接口、上线后调用的大模型等资源费（甲方名下直付，见第六节）",
        ]
    )
    d.table(["项目", "内容"], summary, [3, 13])

    d.h1("二、范围与模块报价")
    d.para("人天为剑来实际投入估算，报价按模块列出，便于甲方按预算取舍。", indent=True)
    rows = [
        [
            str(i),
            it.get("name", ""),
            it.get("content", ""),
            str(it.get("days") or ""),
            money(Decimal(str(it.get("amount") or 0))),
        ]
        for i, it in enumerate(items, 1)
    ]
    rows.append(["", "开发服务费小计", "", f"{days.normalize():f} 人天", money(dev)])
    if ai_fee:
        rows.append(
            [
                "",
                "开发期 AI 费用（包干）",
                f"AI 编程工具与模型 API 用量，{weeks} 周包干，超出由剑来承担",
                "",
                money(ai_fee),
            ]
        )
    rows.append(["", "合计", "", "", money(total)])
    d.table(
        ["#", "模块", "内容", "人天", "报价（元）"], rows, [0.8, 3.2, 8.6, 1.6, 1.8], bold_last=True
    )

    d.h2("上线后运维（可选，质保期满后按月签）")
    d.table(
        ["档位", "包含", "价格"],
        [
            ["基础运维", "缺陷修复、服务器与运行监控；工作日 1 小时内响应", f"{m_low} 元 / 月"],
            [
                "标准运维",
                "基础运维全部内容 + 每月不超过 3 人天的小需求（改文案、调筛选、加字段、调提示词等）+ 每月数据报表",
                f"{m_high} 元 / 月",
            ],
        ],
        [2.6, 10.8, 2.6],
    )
    d.para(
        "运维只负责缺陷修复和小需求。小程序、App、新业务模块等较大功能，另行签订项目合同，按本估算书同一方法报价。",
        indent=True,
    )
    if (q.get("out_of_scope") or "").strip():
        d.h2("明确不在本期")
        d.para(q["out_of_scope"].strip(), indent=True)

    d.h1("三、与常规软件外包对比")
    d.para(
        f"常规外包按长沙及周边市场常见做法估算：项目经理 + UI + 前端 + 后端 + 测试的标准团队，单价 {MARKET_RATE[0]:,}–{MARKET_RATE[1]:,} 元 / 人天，"
        "另计项目管理与商务成本。以下为估算区间，不代表某一家公司的具体报价。",
        indent=True,
        gray=True,
    )
    dd = float(days)
    mk_low = dd * MARKET_DAYS[0] * MARKET_RATE[0]
    mk_high = dd * MARKET_DAYS[1] * MARKET_RATE[1]
    ratio = 1 - float(total) / mk_low if mk_low else 0
    d.table(
        ["维度", "常规软件外包", "剑来科技"],
        [
            [
                "报价",
                f"约 {_round_wan(mk_low)}–{_round_wan(mk_high)} 万元",
                f"{money(total)} 元（比常规低约 {max(ratio, 0) * 100:.0f}% 以上）",
            ],
            [
                "投入人天",
                f"约 {dd * MARKET_DAYS[0]:.0f}–{dd * MARKET_DAYS[1]:.0f} 人天（多角色交接、会议与文档成本高）",
                f"约 {dd:.0f} 人天（AI 辅助编码与测试，同一组人从设计做到上线）",
            ],
            [
                "工期",
                f"{weeks * MARKET_WEEKS[0]:.0f}–{weeks * MARKET_WEEKS[1]:.0f} 周",
                f"{weeks} 周，中途有可演示版本",
            ],
            [
                "沟通方式",
                "销售 → 项目经理 → 开发，需求层层转述；多为线上周报",
                "道长（技术负责人）直接对接；"
                + (f"每周 {onsite} 次到甲方现场，当面演示、当场改；" if onsite else "")
                + "线上 1 小时内响应",
            ],
            [
                "需求变化",
                "改需求走变更单，周期长、加价高",
                "本期内小调整不另收费；大功能另签项目，按同一方法透明报价",
            ],
            [
                "AI 能力",
                "多为外采 SaaS 或简单套壳",
                "剑来日常以 AI 驱动开发，提示词与模型调用按场景调优",
            ],
            [
                "源码与绑定",
                "常见：框架加密、私有组件、服务器在乙方名下",
                "完整源码交付、无加密；服务器与第三方账号开在甲方名下，可独立运营与迁移",
            ],
            ["上线之后", "按次报价，响应慢", f"可选月度运维 {m_low}–{m_high} 元"],
        ],
        [2.6, 6.4, 7.0],
    )
    d.h2("为什么能便宜，又不降质量")
    if onsite:
        d.bullet(
            f"人在现场：每周 {onsite} 次驻场沟通，需求当面讲清、问题当场改，避免做完才发现理解偏差、返工拖期。"
        )
    d.bullet(
        "人力结构不同：剑修是道长亲自带出来的技术驻场，不养销售与多层项目管理，钱花在写代码的人身上。"
    )
    d.bullet("AI 是剑锋：编码、测试、文档大量由 AI 辅助完成，同样的活用更少的人天。")
    d.bullet("工序不砍：设计、评审、测试、部署、培训一项不少，验收按功能清单逐项过。")
    d.bullet("自有底座：登录、权限、后台、埋点等通用能力在剑来已有积累上定制，不从零写。")

    d.h1("四、工期与里程碑")
    if onsite:
        d.para(
            f"开发期间每周 {onsite} 次到甲方现场沟通：对齐本周需求与设计、演示进度、当场确认与调整。",
            indent=True,
        )
    if schedule:
        d.table(
            ["时间", "里程碑", "甲方配合"],
            [[s.get("when", ""), s.get("what", ""), s.get("client", "")] for s in schedule],
            [2.4, 8.2, 5.4],
        )

    d.h1("五、付款方式")
    if pay_plan:
        pay_rows = []
        acc = Decimal(0)
        for i, p in enumerate(pay_plan):
            amt = (
                (total * Decimal(str(p["pct"])) / 100).quantize(Decimal(1))
                if i < len(pay_plan) - 1
                else total - acc
            )
            acc += amt
            pay_rows.append(
                [p.get("name", ""), f"{pct(p['pct'])}%", money(amt), p.get("condition", "")]
            )
        pay_rows.append(["合计", "100%", money(total), ""])
        d.table(
            ["期次", "比例", "金额（元）", "支付条件"],
            pay_rows,
            [2.2, 1.6, 2.4, 9.8],
            bold_last=True,
        )
    d.para(
        f"收款账户（对公）：{PARTY_B['name']}｜{PARTY_B['bank']}｜{PARTY_B['account']}。"
        "价格为含税价，每期收款后开具增值税发票。",
        indent=True,
    )

    d.h1("六、AI 开发费用与第三方资源费")
    if ai_fee:
        d.h2("1. 开发期 AI 费用（已含在总价内，包干）")
        d.para(
            f"剑来以 AI 驱动开发，编码、测试、文档生成中调用的大模型按实际用量计费，这是工期和报价能比常规外包低的原因。"
            f"本期按 {weeks} 周包干 {money(ai_fee)} 元，已含在总价内；超出部分由剑来承担。"
            "因甲方需求变更导致工期延长的，延长期间的 AI 费用另行协商。",
            indent=True,
        )
    d.h2(("2. " if ai_fee else "") + "项目运行资源（甲方承担，以甲方名义开户直付）")
    d.para(
        "服务器、数据库、域名与备案、短信、第三方接口、上线后产品调用的大模型 API 等，以甲方名义开通、甲方直接付费，"
        "剑来凭授权使用和维护，不加价。确需剑来代付的，甲方预付备用金，按实扣减、每月提供账单、结项多退少补。",
        indent=True,
    )

    d.h1("七、报价前提")
    d.bullet(
        "验收以双方确认的《功能清单》为唯一依据；提交验收后 7 个工作日内未书面提出问题的，视为验收通过。"
    )
    d.bullet("清单外的新功能，先报价、经甲方书面确认后开发。")
    d.bullet(
        "源码、设计版权在甲方付清全款后归甲方；剑来原有的通用框架与工具保留权利，甲方在本项目范围内永久免费使用。"
    )
    if onsite:
        d.bullet(
            "驻场沟通每次约 2 小时，地点在甲方长沙办公场所，时间双方提前约定；遇节假日或甲方不便可改为线上，次数不累计。"
        )
    d.bullet("剑来不将本项目转包给第三方；参与开发的驻场人员均为剑来自有团队。")
    for line in (q.get("assumptions") or "").splitlines():
        if line.strip():
            d.bullet(line.strip().lstrip("·-• "))

    d.para()
    d.para(PARTY_B["name"], 11, True, WD_ALIGN_PARAGRAPH.RIGHT)
    d.para(cn_date(made), 10.5, align=WD_ALIGN_PARAGRAPH.RIGHT)

    buffer = BytesIO()
    d.doc.save(buffer)
    return buffer.getvalue()
