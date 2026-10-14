# ruff: noqa: ISC004
"""把合同数据排成 Word。派遣制 = 驻场技术服务（按人月）；项目制先出同一骨架，条款待补。"""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

PARTY_B = {
    "name": "长沙市望城区剑来科技有限责任公司",
    "credit_code": "91430112MAKP5D3E2T",
    "rep": "赵书剑",
    "address": "湖南省长沙市望城区大泽湖街道大泽湖海归小镇研发中心1期2栋102室",
    # 基本存款账户（2026-10-09 开户），合同收款只认这一个户
    "bank": "招商银行股份有限公司长沙学士支行",
    "account": "731914278910001",
}

_DIGITS = "零壹贰叁肆伍陆柒捌玖"
_UNITS = ["", "拾", "佰", "仟"]
_SECTIONS = ["", "万", "亿"]


def rmb_upper(amount: Decimal) -> str:
    """金额大写。合同金额都是整元到分，够用。"""
    cents = int((amount * 100).to_integral_value())
    yuan, jiao, fen = cents // 100, cents // 10 % 10, cents % 10
    if yuan == 0:
        head = "零元"
    else:
        parts: list[str] = []
        section = 0
        need_zero = False
        while yuan > 0:
            chunk = yuan % 10000
            text = ""
            zero = False
            for i in range(4):
                digit = chunk // (10**i) % 10
                if digit == 0:
                    zero = bool(text)
                else:
                    text = _DIGITS[digit] + _UNITS[i] + ("零" if zero else "") + text
                    zero = False
            if chunk:
                tail = "零" if need_zero and parts else ""
                parts.insert(0, text + _SECTIONS[section] + tail)
            need_zero = 0 < chunk < 1000
            yuan //= 10000
            section += 1
        head = "".join(parts) + "元"
    if jiao == 0 and fen == 0:
        return head + "整"
    tail = (_DIGITS[jiao] + "角") if jiao else ("零" if fen else "")
    return head + tail + ((_DIGITS[fen] + "分") if fen else "")


def money(amount: Decimal) -> str:
    return f"{amount:,.2f}".removesuffix(".00")


def cn_date(day: date | None) -> str:
    return f"{day.year} 年 {day.month} 月 {day.day} 日" if day else "    年    月    日"


def totals(staff_fee: Decimal, consult_fee: Decimal, headcount: int, months: int) -> dict:
    monthly = staff_fee * headcount + consult_fee
    return {
        "staff_monthly": staff_fee * headcount,
        "monthly": monthly,
        "total": monthly * months,
    }


def _font(run, size: float = 11, bold: bool = False) -> None:
    run.font.name = "宋体"
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def _para(doc, text: str = "", *, bold: bool = False, size: float = 11, align=None, indent: bool = False):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.line_spacing = 1.5
    if indent:
        para.paragraph_format.first_line_indent = Pt(size * 2)
    if align is not None:
        para.alignment = align
    if text:
        _font(para.add_run(text), size, bold)
    return para


def _cell(cell, text: str, bold: bool = False) -> None:
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(para.add_run(text), 10.5, bold)


def _blank(value: str | None, width: int = 16) -> str:
    return value.strip() if value and value.strip() else "_" * width


def build_contract_docx(c: dict, staff: list[dict]) -> bytes:
    dispatch = c["mode"] == "dispatch"
    sums = totals(c["staff_fee"], c["consult_fee"], len(staff), c["months"])
    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(2.4)
        section.left_margin = section.right_margin = Cm(2.6)

    title = "驻场技术服务合同（按人月）" if dispatch else "技术开发服务合同（项目制）"
    _para(doc, title, bold=True, size=18, align=WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc, f"合同编号：{c['code']}", size=10.5, align=WD_ALIGN_PARAGRAPH.RIGHT)

    _para(doc, f"甲方（委托方）：{_blank(c['party_a'], 24)}", bold=True)
    _para(doc, f"法定代表人 / 负责人：{_blank(c.get('party_a_rep'))}")
    _para(doc, f"地址：{_blank(c.get('party_a_address'), 30)}")
    _para(doc, f"联系人及电话：{_blank(c.get('party_a_contact'), 24)}")
    _para(doc)
    _para(doc, f"乙方（服务方）：{PARTY_B['name']}", bold=True)
    _para(doc, f"统一社会信用代码：{PARTY_B['credit_code']}")
    _para(doc, f"法定代表人：{PARTY_B['rep']}")
    _para(doc, f"地址：{PARTY_B['address']}")
    _para(doc)
    _para(
        doc,
        "甲乙双方根据《中华人民共和国民法典》及相关法律法规，本着平等自愿、诚实信用的原则，"
        "就乙方向甲方提供驻场技术服务事宜协商一致，订立本合同，共同遵守。",
        indent=True,
    )

    n = [0]

    def clause(heading: str, items: list[str]) -> None:
        n[0] += 1
        _para(doc, f"第{_ORDINALS[n[0]]}条　{heading}", bold=True, size=12)
        for i, item in enumerate(items, 1):
            _para(doc, f"{n[0]}.{i}　{item}", indent=True)

    clause(
        "服务内容",
        [
            "乙方指派驻场人员到甲方指定场所，在甲方现有系统和数据基础上，提供系统打通、"
            "软件开发、智能体定制及相关技术服务。具体工作事项由双方在服务期内以书面"
            "（含微信、邮件）方式确认。",
            "乙方指派一名顾问（道长）全程跟进，负责需求判断、方案把关、人员调度和交付质量。",
            f"服务地点：{_blank(c.get('work_place'), 20)}。",
        ],
    )

    n[0] += 1
    _para(doc, f"第{_ORDINALS[n[0]]}条　驻场人员", bold=True, size=12)
    _para(doc, f"{n[0]}.1　乙方指派以下 {len(staff)} 名驻场人员：", indent=True)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, ["序号", "姓名", "岗位", "月服务费（元）"]):
        _cell(cell, text, bold=True)
    for i, person in enumerate(staff, 1):
        row = table.add_row().cells
        _cell(row[0], str(i))
        _cell(row[1], person["legal_name"])
        _cell(row[2], person["role"])
        _cell(row[3], money(c["staff_fee"]))
    _para(doc)
    _para(
        doc,
        f"{n[0]}.2　驻场人员由乙方管理，与乙方存在合作或用工关系，与甲方不构成劳动关系、"
        "劳务关系或劳务派遣关系。驻场人员的报酬、保险及管理责任由乙方承担。",
        indent=True,
    )
    _para(
        doc,
        f"{n[0]}.3　乙方需更换驻场人员的，应提前 3 个工作日通知甲方并做好交接；甲方认为某位"
        "驻场人员不能胜任的，可书面提出，乙方应在 5 个工作日内调整。",
        indent=True,
    )

    clause(
        "服务期限",
        [
            f"服务期共 {c['months']} 个月，自 {cn_date(c['start_on'])} 起至 {cn_date(c['end_on'])} 止。",
            "服务期满前 7 日，任何一方均未书面（含微信、邮件）提出不再续期的，本合同按原人员、"
            "原单价和原条款自动续期一个月，此后每月依此类推。续期期间调整人数或单价的，"
            "以双方书面确认为准。",
            "任何一方不再续期的，本合同在当期服务期满时终止，双方结清费用并完成交接。",
        ],
    )

    fee_lines = [
        f"驻场人员服务费：每人每月人民币 {money(c['staff_fee'])} 元，共 {len(staff)} 人，"
        f"每月合计人民币 {money(sums['staff_monthly'])} 元。",
        f"顾问费：每月人民币 {money(c['consult_fee'])} 元。",
        f"以上两项每月合计人民币 {money(sums['monthly'])} 元（大写：{rmb_upper(sums['monthly'])}）；"
        f"服务期 {c['months']} 个月，合同总价人民币 {money(sums['total'])} 元"
        f"（大写：{rmb_upper(sums['total'])}）。",
        "支付方式："
        + (c.get("pay_terms") or "").strip().rstrip("。")
        + "。甲方以银行转账方式支付至乙方下列对公账户，乙方不接受向任何个人账户付款：",
        f"户名：{PARTY_B['name']}；开户行：{PARTY_B['bank']}；账号：{PARTY_B['account']}。",
        "乙方在每期收款后 5 个工作日内向甲方开具增值税发票。上述价格均为含税价，"
        "按开票时国家税收政策适用的征收率开具（乙方现为增值税小规模纳税人）；"
        "甲方需要增值税专用发票的，应在付款前书面告知，合同价格不因发票类型或税收政策调整而变动。",
        "甲方对当月工作有异议的，应在该服务月结束后 5 个工作日内书面提出；逾期未提出的，"
        "视为甲方确认当月服务。",
        "驻场期间因甲方要求产生的异地差旅、食宿费用，由甲方另行承担，事前经甲方确认。",
    ]
    clause("服务费用与支付", fee_lines)

    clause(
        "甲方的权利与义务",
        [
            "指定一名对接人，负责需求确认、进度沟通和成果确认。",
            "提供驻场所需的办公位、网络、系统账号、数据访问权限及必要的业务说明。",
            "按本合同约定按时支付费用。",
            "未经乙方书面同意，服务期内及期满后 12 个月内，不直接或间接聘用乙方驻场人员。",
        ],
    )
    clause(
        "乙方的权利与义务",
        [
            "按约定指派驻场人员，并保证其具备完成工作所需的能力。",
            "驻场人员遵守甲方合理的现场管理制度和信息安全要求。",
            "线上问题在一小时内响应；每周向甲方对接人同步一次进度、问题和下周安排。",
            "服务期满或合同终止时，交还甲方提供的账号、权限和资料，并完成工作交接。",
        ],
    )
    clause(
        "知识产权",
        [
            "乙方在本合同履行过程中为甲方形成的全部成果（含源代码、文档、设计、模型、配置等），"
            "其知识产权自甲方付清对应费用之日起归甲方所有。乙方应按甲方要求完整交付源代码及相关文档。",
            "乙方在本合同之前已有的通用工具、框架和方法，知识产权仍归乙方；"
            "其中被用于上述成果的部分，甲方在自身业务范围内享有永久、免费的使用权。",
            "甲方提供的业务数据、资料及甲方原有系统，其权利归甲方，乙方仅为履行本合同使用。",
        ],
    )
    clause(
        "保密与数据安全",
        [
            "双方对履行本合同知悉的对方商业秘密、技术资料和业务数据负有保密义务，"
            "保密期限为本合同终止后 2 年。",
            "乙方及驻场人员不得将甲方业务数据存入个人设备或带出甲方环境；确需带出的，"
            "须事先取得甲方书面同意。",
            "未经甲方书面同意，乙方不得在对外宣传中使用甲方名称；可以只写行业。",
        ],
    )
    clause(
        "违约责任",
        [
            "甲方逾期付款的，每逾期一日按应付未付金额的万分之五支付违约金；首期款未按时到账的，"
            "乙方进场日相应顺延，服务期不因此缩短；其他款项逾期超过 15 日的，乙方有权暂停服务。",
            "乙方驻场人员无正当理由连续 3 个工作日未到岗，且乙方未补派的，甲方有权按缺勤天数"
            "相应扣减当月服务费。",
            "任何一方违反保密义务给对方造成损失的，应赔偿对方的实际损失。",
            "除故意或重大过失外，乙方在本合同项下的累计赔偿责任以甲方已支付的服务费总额为限。",
        ],
    )
    clause(
        "合同变更与解除",
        [
            "变更服务内容、人数或期限的，由双方书面确认，费用相应调整。",
            "任何一方可提前 7 日书面通知对方解除本合同，费用按实际服务天数结算，多退少补。",
        ],
    )
    clause(
        "争议解决",
        ["因本合同产生的争议，双方先友好协商；协商不成的，提交乙方所在地有管辖权的人民法院诉讼解决。"],
    )
    other = ["本合同一式两份，甲乙双方各执一份，自双方签字或盖章之日起生效。"]
    other.append("本合同未尽事宜，由双方另行签订补充协议，补充协议与本合同具有同等效力。")
    if (c.get("extra_terms") or "").strip():
        other.append("补充约定：" + c["extra_terms"].strip())
    clause("其他", other)

    _para(doc)
    sign = doc.add_table(rows=4, cols=2)
    rows = [
        ("甲方（盖章）：", "乙方（盖章）："),
        ("", ""),
        ("代表人（签字）：", "代表人（签字）："),
        ("日期：    年    月    日", "日期：    年    月    日"),
    ]
    for row, (left, right) in zip(sign.rows, rows):
        for cell, text in zip(row.cells, (left, right)):
            _font(cell.paragraphs[0].add_run(text), 11)
            cell.paragraphs[0].paragraph_format.space_after = Pt(14)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


_ORDINALS = ["", "一", "二", "三", "四", "五", "六", "七", "八", "九", "十", "十一", "十二", "十三"]
